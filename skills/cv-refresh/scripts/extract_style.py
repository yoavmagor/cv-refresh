#!/usr/bin/env python3
"""Detect the visual style of an existing CV so it can be mirrored.

Usage:
  python3 extract_style.py INPUT --out style.json

Supports DOCX and PDF (Markdown/TXT inputs get sensible defaults). The result
feeds `render_cv.py --template mirror --style style.json`. Detection is
heuristic: fonts, sizes, accent color, margins, page size, header alignment,
heading case/rule and section order. Multi-column or graphic-heavy designs are
flagged and rendered as a clean single column (better for applicant tracking
systems anyway).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import emit, fail  # noqa: E402
from sections import canonical_section  # noqa: E402

DEFAULT_STYLE = {
    "page_size": "letter",
    "margins_in": {"top": 0.7, "bottom": 0.7, "left": 0.75, "right": 0.75},
    "font_body": "Calibri",
    "font_heading": "Calibri",
    "size_body": 10.5,
    "size_name": 20,
    "size_heading": 11.5,
    "accent_color": "1F3864",
    "heading_color": "1F3864",
    "heading_case": "upper",
    "heading_bold": True,
    "heading_rule": True,
    "header_alignment": "left",
    "section_order": [],
    "columns": 1,
    "notes": [],
}

PS_FONT_MAP = {
    "arialmt": "Arial", "arial": "Arial", "helvetica": "Arial", "helveticaneue": "Helvetica Neue",
    "timesnewromanpsmt": "Times New Roman", "timesnewroman": "Times New Roman", "times": "Times New Roman",
    "timesroman": "Times New Roman", "couriernewpsmt": "Courier New", "courier": "Courier New",
    "carlito": "Calibri", "caladea": "Cambria", "liberationsans": "Arial",
    "liberationserif": "Times New Roman", "dejavusans": "Arial", "dejavuserif": "Georgia",
    "calibri": "Calibri", "cambria": "Cambria", "georgia": "Georgia", "garamond": "Garamond",
    "ebgaramond": "EB Garamond", "verdana": "Verdana", "tahoma": "Tahoma", "segoeui": "Segoe UI",
    "opensans": "Open Sans", "roboto": "Roboto", "lato": "Lato", "sourcesanspro": "Source Sans Pro",
    "montserrat": "Montserrat", "aptos": "Aptos", "bookantiqua": "Book Antiqua",
    "palatinolinotype": "Palatino Linotype", "trebuchetms": "Trebuchet MS",
}
STYLE_SUFFIX = re.compile(r"[-,](bold|italic|oblique|light|medium|semibold|black|regular|"
                          r"boldmt|italicmt|bolditalic|bolditalicmt|mt|ps|psmt|condensed).*$", re.I)


def normalize_font(raw: str | None) -> str | None:
    if not raw:
        return None
    name = raw.split("+", 1)[-1]           # drop subset prefix "ABCDEF+"
    name = STYLE_SUFFIX.sub("", name)
    key = re.sub(r"[^a-z]", "", name.lower())
    key = re.sub(r"(bold|italic|regular|mt|psmt)$", "", key)
    if key in PS_FONT_MAP:
        return PS_FONT_MAP[key]
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name).strip() or None


def page_size_name(width_pt: float, height_pt: float) -> str:
    return "a4" if abs(width_pt - 595) < 12 and abs(height_pt - 842) < 12 else "letter"


def to_hex(color) -> str | None:
    """pdfplumber colors: gray (1 value), RGB (3), CMYK (4), floats 0..1."""
    if color is None:
        return None
    if isinstance(color, (int, float)):
        color = (color,)
    try:
        vals = [float(v) for v in color]
    except (TypeError, ValueError):
        return None
    if len(vals) == 1:
        r = g = b = vals[0]
    elif len(vals) == 3:
        r, g, b = vals
    elif len(vals) == 4:
        c, m, y, k = vals
        r, g, b = (1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)
    else:
        return None
    return "".join(f"{max(0, min(255, round(v * 255))):02X}" for v in (r, g, b))


def is_neutral(hexcolor: str | None) -> bool:
    if not hexcolor:
        return True
    r, g, b = (int(hexcolor[i:i + 2], 16) for i in (0, 2, 4))
    return max(r, g, b) - min(r, g, b) < 24  # grays, black, white


def heading_case(texts: list[str]) -> str:
    letters = [t for t in texts if re.search(r"[A-Za-z]", t)]
    if letters and sum(t.isupper() for t in letters) >= len(letters) / 2:
        return "upper"
    return "title"


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def style_from_pdf(path: Path) -> dict:
    import pdfplumber  # type: ignore
    style = json.loads(json.dumps(DEFAULT_STYLE))
    with pdfplumber.open(str(path)) as pdf:
        first = pdf.pages[0]
        style["page_size"] = page_size_name(first.width, first.height)
        all_chars = [c for p in pdf.pages for c in p.chars if c["text"].strip()]
        if not all_chars:
            style["notes"].append("No text layer; defaults used.")
            return style
        sizes = Counter(round(c["size"] * 2) / 2 for c in all_chars)
        fonts = Counter(normalize_font(c["fontname"]) for c in all_chars)
        body_size = sizes.most_common(1)[0][0]
        style["size_body"] = max(9.0, min(12.0, body_size))
        style["font_body"] = fonts.most_common(1)[0][0] or style["font_body"]

        xs0 = [c["x0"] for c in first.chars if c["text"].strip()]
        xs1 = [c["x1"] for c in first.chars if c["text"].strip()]
        tops = [c["top"] for c in first.chars if c["text"].strip()]
        pt = 72.0
        style["margins_in"] = {
            "left": round(max(0.4, min(1.2, min(xs0) / pt)), 2),
            "right": round(max(0.4, min(1.2, (first.width - max(xs1)) / pt)), 2),
            "top": round(max(0.4, min(1.2, min(tops) / pt)), 2),
            "bottom": style["margins_in"]["bottom"],
        }

        lines = []
        for page in pdf.pages:
            for line in page.extract_text_lines(x_tolerance=1.5):
                chars = [c for c in line["chars"] if c["text"].strip()]
                if not chars:
                    continue
                lines.append({
                    "page": page.page_number, "text": line["text"].strip(),
                    "x0": line["x0"], "x1": line["x1"], "top": line["top"],
                    "size": max(c["size"] for c in chars),
                    "bold": sum("bold" in c["fontname"].lower() for c in chars) > len(chars) / 2,
                    "font": Counter(normalize_font(c["fontname"]) for c in chars).most_common(1)[0][0],
                    "color": Counter(to_hex(c.get("non_stroking_color")) for c in chars).most_common(1)[0][0],
                    "width": first.width,
                })

        page1 = [ln for ln in lines if ln["page"] == 1]
        if page1:
            name_line = max(page1[:6], key=lambda ln: ln["size"])
            style["size_name"] = round(min(32.0, max(14.0, name_line["size"])), 1)
            centre = (name_line["x0"] + name_line["x1"]) / 2
            style["header_alignment"] = "center" if abs(centre - first.width / 2) < 30 and name_line["x0"] > min(xs0) + 20 else "left"
            if not is_neutral(name_line["color"]):
                style["accent_color"] = name_line["color"]

        headings = [ln for ln in lines if len(ln["text"]) <= 40 and canonical_section(ln["text"])
                    and (ln["size"] > body_size + 0.4 or ln["bold"] or ln["text"].isupper())]
        if headings:
            style["font_heading"] = Counter(h["font"] for h in headings).most_common(1)[0][0] or style["font_body"]
            style["size_heading"] = round(Counter(round(h["size"] * 2) / 2 for h in headings).most_common(1)[0][0], 1)
            style["heading_bold"] = sum(h["bold"] for h in headings) >= len(headings) / 2
            style["heading_case"] = heading_case([h["text"] for h in headings])
            colors = Counter(h["color"] for h in headings)
            top_color = colors.most_common(1)[0][0]
            style["heading_color"] = top_color if top_color and not is_neutral(top_color) else "000000"
            if style["heading_color"] != "000000":
                style["accent_color"] = style["heading_color"]
            seen = []
            for h in headings:
                key = canonical_section(h["text"])
                if key and key not in seen:
                    seen.append(key)
            style["section_order"] = seen
            # horizontal rules drawn near headings
            rules = 0
            for h in headings:
                page = pdf.pages[h["page"] - 1]
                for obj in list(page.lines) + list(page.rects):
                    if abs(obj["top"] - obj["bottom"]) < 2.5 and (obj["x1"] - obj["x0"]) > first.width * 0.4 \
                            and 0 <= obj["top"] - h["top"] < 22:
                        rules += 1
                        break
            style["heading_rule"] = rules >= max(1, len(headings) // 2)
        else:
            style["notes"].append("No section headings recognized; heading style defaults used.")

        right_starts = [ln for ln in page1 if ln["x0"] > first.width * 0.45 and len(ln["text"]) > 25]
        if len(right_starts) >= 6:
            style["columns"] = 2
            style["notes"].append("Original looks multi-column; output is a single column (ATS-friendly).")
        if any(p.images for p in pdf.pages):
            style["notes"].append("Original contains images (photo/logos/icons); they are not carried over.")
    return style


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def style_from_docx(path: Path) -> dict:
    import docx  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    style = json.loads(json.dumps(DEFAULT_STYLE))
    doc = docx.Document(str(path))
    sec = doc.sections[0]
    if sec.page_width and sec.page_height:
        style["page_size"] = page_size_name(sec.page_width.pt, sec.page_height.pt)
    for side in ("top", "bottom", "left", "right"):
        val = getattr(sec, f"{side}_margin")
        if val is not None:
            style["margins_in"][side] = round(max(0.4, min(1.2, val.inches)), 2)

    normal = doc.styles["Normal"]
    default_font = normal.font.name or _theme_minor_font(doc) or "Calibri"
    default_size = normal.font.size.pt if normal.font.size else 11.0

    def run_font(run, para):
        return run.font.name or (para.style.font.name if para.style is not None else None) or default_font

    def run_size(run, para):
        if run.font.size:
            return run.font.size.pt
        if para.style is not None and para.style.font.size:
            return para.style.font.size.pt
        return default_size

    fonts, sizes = Counter(), Counter()
    paras = [p for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(p for p in cell.paragraphs if p.text.strip())
    for p in paras:
        for r in p.runs:
            n = len(r.text.strip())
            if n:
                fonts[run_font(r, p)] += n
                sizes[run_size(r, p)] += n
    if fonts:
        style["font_body"] = normalize_font(fonts.most_common(1)[0][0]) or style["font_body"]
    if sizes:
        style["size_body"] = max(9.0, min(12.0, sizes.most_common(1)[0][0]))

    if paras:
        top = paras[:4]
        name_para = max(top, key=lambda p: max((run_size(r, p) for r in p.runs), default=0))
        style["size_name"] = round(min(32.0, max(14.0, max((run_size(r, name_para) for r in name_para.runs), default=20))), 1)
        align = name_para.alignment if name_para.alignment is not None else (
            name_para.style.paragraph_format.alignment if name_para.style is not None else None)
        style["header_alignment"] = "center" if align is not None and int(align) == 1 else "left"
        color = _run_color(name_para)
        if color and not is_neutral(color):
            style["accent_color"] = color

    headings = [p for p in paras if len(p.text.strip()) <= 40 and canonical_section(p.text.strip())]
    if headings:
        h_fonts = Counter(normalize_font(run_font(r, p)) for p in headings for r in p.runs if r.text.strip())
        h_sizes = Counter(run_size(r, p) for p in headings for r in p.runs if r.text.strip())
        style["font_heading"] = h_fonts.most_common(1)[0][0] if h_fonts else style["font_body"]
        style["size_heading"] = h_sizes.most_common(1)[0][0] if h_sizes else style["size_heading"]
        bold_votes = [any(r.bold for r in p.runs) or bool(p.style is not None and p.style.font.bold) for p in headings]
        style["heading_bold"] = sum(bold_votes) >= len(bold_votes) / 2
        caps_votes = [p.text.isupper() or any(r.font.all_caps for r in p.runs) or
                      bool(p.style is not None and p.style.font.all_caps) for p in headings]
        style["heading_case"] = "upper" if sum(caps_votes) >= len(caps_votes) / 2 else "title"
        colors = Counter(_run_color(p) for p in headings)
        hc = colors.most_common(1)[0][0]
        style["heading_color"] = hc if hc and not is_neutral(hc) else "000000"
        if style["heading_color"] != "000000":
            style["accent_color"] = style["heading_color"]
        rule_votes = [_has_bottom_border(p, qn) for p in headings]
        style["heading_rule"] = sum(rule_votes) >= len(rule_votes) / 2
        order = []
        for p in headings:
            key = canonical_section(p.text.strip())
            if key not in order:
                order.append(key)
        style["section_order"] = order
    else:
        style["notes"].append("No section headings recognized; heading style defaults used.")
    if doc.tables:
        style["notes"].append("Original uses tables for layout; output is a single column (ATS-friendly).")
    if doc.inline_shapes and len(doc.inline_shapes):
        style["notes"].append("Original contains images; they are not carried over.")
    return style


def _theme_minor_font(doc) -> str | None:
    try:
        styles_xml = doc.styles.element.xml
        m = re.search(r'<w:rFonts[^>]*w:ascii="([^"]+)"', styles_xml)
        return m.group(1) if m else None
    except Exception:
        return None


def _run_color(para) -> str | None:
    for r in para.runs:
        if r.text.strip() and r.font.color is not None and r.font.color.type is not None:
            try:
                return str(r.font.color.rgb)
            except Exception:
                return None
    if para.style is not None and para.style.font.color is not None and para.style.font.color.type is not None:
        try:
            return str(para.style.font.color.rgb)
        except Exception:
            return None
    return None


def _has_bottom_border(para, qn) -> bool:
    ppr = para._p.pPr
    for node in (ppr, getattr(para.style, "element", None)):
        if node is not None and node.find(".//" + qn("w:pBdr")) is not None:
            return node.find(".//" + qn("w:bottom")) is not None
    return False


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    src = Path(args.input)
    if not src.exists():
        fail(f"File not found: {src}")
    suffix = src.suffix.lower()
    try:
        if suffix == ".pdf":
            style = style_from_pdf(src)
        elif suffix == ".docx":
            style = style_from_docx(src)
        else:
            style = json.loads(json.dumps(DEFAULT_STYLE))
            style["notes"].append(f"No visual style in a {suffix} file; defaults used.")
    except ImportError as exc:
        fail(f"Missing library: {exc.name}. Run scripts/doctor.py for install hints.")
    style["source_file"] = str(src)
    Path(args.out).write_text(json.dumps(style, indent=2) + "\n", encoding="utf-8")
    emit({"ok": True, "style_file": args.out, "style": style})


if __name__ == "__main__":
    main()
