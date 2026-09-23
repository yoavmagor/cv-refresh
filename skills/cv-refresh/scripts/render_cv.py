#!/usr/bin/env python3
"""Render a cv.json into DOCX (+ PDF when LibreOffice is available).

Usage:
  python3 render_cv.py --cv new.cv.json --out-dir OUT --basename Jane_Doe_CV \
      --template mirror --style style.json --max-pages 2 --autofit [--previews]

  # layout options for the user to choose from (one file per template + page-1 previews)
  python3 render_cv.py --cv new.cv.json --out-dir OUT/options --basename Jane_Doe_CV \
      --template classic,modern,compact,mirror --style style.json --max-pages 2 --autofit

Templates:
  mirror   fonts, sizes, colors, margins and heading style detected from the
           user's original CV (style.json); always a single column
  classic  serif (Cambria), centered header, black headings with a rule
  modern   sans (Calibri), left header, colored headings
  compact  sans (Arial), tight margins and spacing, for dense 2-page CVs

--autofit only changes typography (spacing, font size, margins) within
readable limits. If the CV still exceeds --max-pages, the JSON output says so
and the content has to be trimmed (see references/writing-guide.md).
Nothing here adds comments, highlights or placeholder text to the document.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (docx_to_pdf, emit, fail, find_soffice, last_page_fill,  # noqa: E402
                    load_json, pdf_page_count, render_previews)

try:
    from docx import Document  # type: ignore
    from docx.enum.style import WD_STYLE_TYPE  # type: ignore
    from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # type: ignore
    from docx.opc.constants import RELATIONSHIP_TYPE as RT  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Inches, Pt, RGBColor  # type: ignore
except ImportError:  # pragma: no cover
    Document = None

sys.path.insert(0, str(Path(__file__).resolve().parent))
from extract_style import DEFAULT_STYLE  # noqa: E402

TEMPLATES = {
    "classic": {
        "font_body": "Cambria", "font_heading": "Cambria", "size_body": 10.5, "size_name": 22,
        "size_heading": 11.5, "accent_color": "000000", "heading_color": "000000",
        "heading_case": "upper", "heading_bold": True, "heading_rule": True,
        "header_alignment": "center",
        "margins_in": {"top": 0.7, "bottom": 0.7, "left": 0.8, "right": 0.8},
    },
    "modern": {
        "font_body": "Calibri", "font_heading": "Calibri", "size_body": 10.5, "size_name": 24,
        "size_heading": 12, "accent_color": "1F4E79", "heading_color": "1F4E79",
        "heading_case": "upper", "heading_bold": True, "heading_rule": False,
        "header_alignment": "left",
        "margins_in": {"top": 0.65, "bottom": 0.65, "left": 0.75, "right": 0.75},
    },
    "compact": {
        "font_body": "Arial", "font_heading": "Arial", "size_body": 9.5, "size_name": 18,
        "size_heading": 10.5, "accent_color": "222222", "heading_color": "000000",
        "heading_case": "upper", "heading_bold": True, "heading_rule": True,
        "header_alignment": "left",
        "margins_in": {"top": 0.5, "bottom": 0.5, "left": 0.55, "right": 0.55},
    },
}

DEFAULT_TITLES = {
    "summary": "Summary", "experience": "Experience", "skills": "Skills",
    "education": "Education", "certifications": "Certifications", "projects": "Projects",
    "languages": "Languages",
}
DEFAULT_ORDER = ["summary", "experience", "skills", "projects", "education",
                 "certifications", "languages"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

AUTOFIT_STEPS = [
    ("tighten spacing", {"spacing": 0.8}),
    ("body font -0.5pt", {"size_delta": -0.5}),
    ("margins -0.1in", {"margin_delta": -0.1}),
    ("tighten spacing more", {"spacing": 0.6}),
    ("body font -0.5pt", {"size_delta": -0.5}),
    ("margins -0.1in", {"margin_delta": -0.1}),
]
MIN_BODY_PT = 9.0
MIN_MARGIN_IN = 0.45


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def fmt_date(value) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.lower() in ("present", "current", "now", "today"):
        return "Present"
    m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-\d{1,2})?", s)
    if m:
        month = int(m.group(2))
        return f"{MONTHS[month - 1]} {m.group(1)}" if 1 <= month <= 12 else m.group(1)
    return s


def date_range(start, end) -> str:
    a, b = fmt_date(start), fmt_date(end)
    if a and b:
        return a if a == b else f"{a} \u2013 {b}"
    return a or b


def set_font(element_holder, name: str) -> None:
    """Set a font on a style or run and drop theme-font attributes that override it."""
    rpr = element_holder.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def bottom_border(paragraph_format_holder, color: str) -> None:
    ppr = paragraph_format_holder.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def add_hyperlink(paragraph, url: str, text: str, color: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def normalize_url(value: str) -> str:
    if value.startswith(("http://", "https://", "mailto:")):
        return value
    if "@" in value and " " not in value:
        return f"mailto:{value}"
    return f"https://{value}"


# --------------------------------------------------------------------------
# document builder
# --------------------------------------------------------------------------

class Builder:
    def __init__(self, cv: dict, p: dict):
        self.cv = cv
        self.p = p
        self.sp = p.get("spacing", 1.0)
        self.doc = Document()
        self._setup_page()
        self._setup_styles()

    # page & styles ---------------------------------------------------------
    def _setup_page(self):
        sec = self.doc.sections[0]
        if self.p.get("page_size") == "a4":
            sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        m = self.p["margins_in"]
        sec.top_margin, sec.bottom_margin = Inches(m["top"]), Inches(m["bottom"])
        sec.left_margin, sec.right_margin = Inches(m["left"]), Inches(m["right"])
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin

    def _para_style(self, name, size, bold=False, italic=False, color=None, font=None,
                    before=0.0, after=0.0, align=None, keep_next=False):
        styles = self.doc.styles
        st = styles[name] if name in [s.name for s in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != "Normal":
            st.base_style = styles["Normal"]
        st.quick_style = True
        set_font(st.element, font or self.p["font_body"])
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = italic
        if color:
            st.font.color.rgb = RGBColor.from_string(color)
        pf = st.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.0
        pf.keep_with_next = keep_next
        if align is not None:
            pf.alignment = align
        return st

    def _setup_styles(self):
        p, sp = self.p, self.sp
        body = p["size_body"]
        align = WD_ALIGN_PARAGRAPH.CENTER if p["header_alignment"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
        self._para_style("Normal", body, color="000000", after=0)
        self._para_style("CV Name", p["size_name"], bold=True, color=p["accent_color"],
                         font=p["font_heading"], after=1, align=align)
        self._para_style("CV Headline", body + 1.5, color="333333", after=2, align=align)
        self._para_style("CV Contact", body - 0.5, color="333333", after=4 * sp, align=align)
        heading = self._para_style("CV Heading", p["size_heading"], bold=p["heading_bold"],
                                   color=p["heading_color"], font=p["font_heading"],
                                   before=10 * sp, after=3 * sp, keep_next=True)
        heading.font.all_caps = p["heading_case"] == "upper"
        if p["heading_rule"]:
            bottom_border(heading.element, p["heading_color"] if p["heading_color"] != "000000" else "444444")
        self._para_style("CV Role", body + 0.5, before=6 * sp, after=0, keep_next=True)
        self._para_style("CV Meta", body - 0.5, italic=True, color="555555", after=1.5 * sp, keep_next=True)
        self._para_style("CV Positions", body, color="222222", after=1 * sp, keep_next=True)
        self._para_style("CV Body", body, after=3 * sp)
        self._para_style("CV Tech", body - 0.5, color="444444", before=1 * sp, after=1 * sp)
        lb = self.doc.styles["List Bullet"]
        set_font(lb.element, p["font_body"])
        lb.font.size = Pt(body)
        lb.paragraph_format.space_after = Pt(1.5 * sp)
        lb.paragraph_format.line_spacing = 1.0
        lb.paragraph_format.left_indent = Inches(0.22)
        lb.paragraph_format.first_line_indent = Inches(-0.16)

    # primitives ------------------------------------------------------------
    def para(self, style, text=""):
        para = self.doc.add_paragraph(style=style)
        if text:
            para.add_run(text)
        return para

    def heading(self, text):
        self.para("CV Heading", text)

    def bullet(self, text):
        self.para("List Bullet", text)

    def role_line(self, left_bold, left_rest, right):
        para = self.para("CV Role")
        para.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        para.add_run(left_bold).bold = True
        if left_rest:
            para.add_run(left_rest)
        if right:
            para.add_run("\t" + right)
        return para

    # sections --------------------------------------------------------------
    def header(self):
        b = self.cv.get("basics", {})
        self.para("CV Name", b.get("name", ""))
        if b.get("headline"):
            self.para("CV Headline", b["headline"])
        items = []
        for key in ("location", "phone", "email"):
            if b.get(key):
                items.append((b[key], key == "email"))
        for link in b.get("links", []) or []:
            label = link.get("display") or re.sub(r"^https?://(www\.)?", "", link.get("url", "")).rstrip("/")
            items.append((label, True, link.get("url")))
        if items:
            para = self.para("CV Contact")
            for i, item in enumerate(items):
                if i:
                    para.add_run("  |  ")
                text, linked = item[0], item[1]
                if linked:
                    add_hyperlink(para, normalize_url(item[2] if len(item) > 2 and item[2] else text),
                                  text, "000000" if self.p["accent_color"] in ("000000", "222222") else self.p["accent_color"])
                else:
                    para.add_run(text)

    def summary(self):
        paras = self.cv.get("summary") or []
        if isinstance(paras, str):
            paras = [paras]
        if not paras:
            return
        self.heading(self.title("summary"))
        for text in paras:
            self.para("CV Body", text)

    def role_block(self, item: dict):
        title = item.get("title", "")
        org = item.get("company") or item.get("organization") or item.get("org") or ""
        positions = item.get("positions") or []
        if positions:
            # One employer, several titles, shared bullets (used when the user didn't
            # say which work belongs to which title -- never guess the attribution).
            starts = [str(p.get("start", "")) for p in positions if p.get("start")]
            ends = [str(p.get("end", "")) for p in positions if p.get("end")]
            end = "present" if any(e.lower() == "present" for e in ends) else (max(ends) if ends else "")
            self.role_line(org, "", date_range(min(starts) if starts else "", end))
            para = self.para("CV Positions")
            for i, pos in enumerate(positions):
                if i:
                    para.add_run("  \u00b7  ")
                para.add_run(pos.get("title", "")).bold = True
                span = date_range(pos.get("start"), pos.get("end"))
                if span:
                    para.add_run(f" ({span})")
        else:
            self.role_line(title, f"  \u00b7  {org}" if org else "", date_range(item.get("start"), item.get("end")))
        meta = " \u00b7 ".join(x for x in (item.get("location"), item.get("company_blurb")) if x)
        if meta:
            self.para("CV Meta", meta)
        if item.get("description"):
            self.para("CV Body", item["description"])
        for text in item.get("bullets", []) or []:
            self.bullet(text)
        if item.get("tech"):
            para = self.para("CV Tech")
            para.add_run("Tech: ").bold = True
            para.add_run(", ".join(item["tech"]))

    def experience(self):
        roles = self.cv.get("experience") or []
        if not roles:
            return
        self.heading(self.title("experience"))
        for role in roles:
            self.role_block(role)

    def skills(self):
        groups = self.cv.get("skills") or []
        if not groups:
            return
        self.heading(self.title("skills"))
        for group in groups:
            para = self.para("CV Body")
            para.paragraph_format.space_after = Pt(1.5 * self.sp)
            if group.get("group"):
                para.add_run(f"{group['group']}: ").bold = True
            para.add_run(", ".join(group.get("items", [])))

    def education(self):
        items = self.cv.get("education") or []
        if not items:
            return
        self.heading(self.title("education"))
        for ed in items:
            deg, field = ed.get("degree") or "", ed.get("field") or ""
            # "B.Sc. Computer Science" but "Bachelor of Science, Computer Science"
            degree = f"{deg} {field}" if deg.endswith(".") and field else ", ".join(x for x in (deg, field) if x)
            inst = ed.get("institution", "")
            self.role_line(degree or inst, f"  \u00b7  {inst}" if degree and inst else "",
                           date_range(ed.get("start"), ed.get("end")))
            for text in ed.get("details", []) or []:
                self.bullet(text)

    def certifications(self):
        items = self.cv.get("certifications") or []
        if not items:
            return
        self.heading(self.title("certifications"))
        for c in items:
            tail = ", ".join(str(x) for x in (c.get("issuer"), c.get("year")) if x)
            para = self.para("CV Body")
            para.paragraph_format.space_after = Pt(1.5 * self.sp)
            para.add_run(c.get("name", "")).bold = True
            if tail:
                para.add_run(f" \u2014 {tail}")

    def projects(self):
        items = self.cv.get("projects") or []
        if not items:
            return
        self.heading(self.title("projects"))
        for pr in items:
            para = self.para("CV Body")
            para.paragraph_format.space_after = Pt(1.5 * self.sp)
            para.add_run(pr.get("name", "")).bold = True
            if pr.get("description"):
                para.add_run(f" \u2014 {pr['description']}")
            if pr.get("link"):
                para.add_run("  ")
                add_hyperlink(para, normalize_url(pr["link"]),
                              re.sub(r"^https?://(www\.)?", "", pr["link"]).rstrip("/"), "555555")
            for text in pr.get("bullets", []) or []:
                self.bullet(text)

    def languages(self):
        items = self.cv.get("languages") or []
        if not items:
            return
        self.heading(self.title("languages"))
        parts = [f"{l.get('language')} ({l['level']})" if l.get("level") else l.get("language", "") for l in items]
        self.para("CV Body", "  \u00b7  ".join(parts))

    def extra(self, section: dict):
        self.heading(section.get("title", ""))
        style = section.get("style", "bullets")
        for item in section.get("items", []) or []:
            if isinstance(item, dict):
                self.role_block(item)
            elif style == "lines":
                para = self.para("CV Body", item)
                para.paragraph_format.space_after = Pt(1.5 * self.sp)
            else:
                self.bullet(item)

    def title(self, key: str) -> str:
        return (self.cv.get("meta", {}).get("section_titles", {}) or {}).get(key, DEFAULT_TITLES.get(key, key.title()))

    def build(self, out_path: Path):
        self.header()
        self.summary()  # the TL;DR always sits directly under the header
        order = self.cv.get("meta", {}).get("section_order") or DEFAULT_ORDER
        extras = {s.get("key") or s.get("title"): s for s in self.cv.get("extra_sections", []) or []}
        done = {"summary"}
        for key in order:
            if key in done:
                continue
            done.add(key)
            if hasattr(self, key) and key in DEFAULT_TITLES:
                getattr(self, key)()
            elif key in extras:
                self.extra(extras[key])
        for key in DEFAULT_ORDER:
            if key not in done and key != "summary":
                getattr(self, key)()
                done.add(key)
        for key, section in extras.items():
            if key not in done:
                self.extra(section)
        core = self.doc.core_properties
        core.title = f"{self.cv.get('basics', {}).get('name', '')} \u2014 CV".strip(" \u2014")
        core.author = self.cv.get("basics", {}).get("name", "")
        core.comments = ""
        self.doc.save(str(out_path))


# --------------------------------------------------------------------------
# parameters, autofit, CLI
# --------------------------------------------------------------------------

def base_params(template: str, style: dict | None) -> dict:
    params = copy.deepcopy(DEFAULT_STYLE)
    if template == "mirror":
        if style:
            params.update({k: v for k, v in style.items() if k in DEFAULT_STYLE})
    else:
        params.update(copy.deepcopy(TEMPLATES[template]))
        if style and style.get("page_size"):
            params["page_size"] = style["page_size"]  # keep the user's paper size
    params["spacing"] = 1.0
    return params


def apply_step(params: dict, step: dict) -> bool:
    """Mutate params; return False if the step can't change anything."""
    changed = False
    if "spacing" in step and step["spacing"] < params.get("spacing", 1.0):
        params["spacing"] = step["spacing"]
        changed = True
    if "size_delta" in step and params["size_body"] + step["size_delta"] >= MIN_BODY_PT:
        params["size_body"] += step["size_delta"]
        params["size_heading"] = max(params["size_body"] + 0.5, params["size_heading"] + step["size_delta"])
        params["size_name"] = max(16, params["size_name"] - 1)
        changed = True
    if "margin_delta" in step:
        for side, val in params["margins_in"].items():
            new = round(val + step["margin_delta"], 2)
            if new >= MIN_MARGIN_IN and new < val:
                params["margins_in"][side] = new
                changed = True
    return changed


def render_one(cv, template, style, out_dir: Path, basename: str, max_pages: int, autofit: bool,
               previews: bool, first_page_preview: bool) -> dict:
    params = base_params(template, style)
    docx_path = out_dir / f"{basename}.docx"
    steps_applied: list[str] = []
    pdf_path = pages = None
    step_iter = iter(AUTOFIT_STEPS)
    while True:
        Builder(cv, params).build(docx_path)
        pdf_path = docx_to_pdf(docx_path, out_dir)
        pages = pdf_page_count(pdf_path) if pdf_path else None
        if pages is None or pages <= max_pages or not autofit:
            break
        for label, step in step_iter:
            if apply_step(params, step):
                steps_applied.append(label)
                break
        else:
            break
    # A last page that is nearly empty looks unfinished. Try to pull it back
    # with the same typographic steps; keep the result only if it works.
    fill = last_page_fill(pdf_path) if pdf_path else None
    if autofit and pages and pages > 1 and fill is not None and fill < 0.2:
        trial = copy.deepcopy(params)
        trial_steps: list[str] = []
        for label, step in AUTOFIT_STEPS:
            if not apply_step(trial, step):
                continue
            trial_steps.append(label)
            Builder(cv, trial).build(docx_path)
            trial_pdf = docx_to_pdf(docx_path, out_dir)
            trial_pages = pdf_page_count(trial_pdf) if trial_pdf else None
            if trial_pages is not None and trial_pages < pages:
                params, pages, pdf_path = trial, trial_pages, trial_pdf
                steps_applied += [f"{s} (to remove a near-empty last page)" for s in trial_steps]
                break
        else:
            Builder(cv, params).build(docx_path)  # restore the untouched version
            pdf_path = docx_to_pdf(docx_path, out_dir)
            steps_applied.append("could not remove the near-empty last page by typography alone; trim or add content")
    result = {
        "template": template,
        "docx": str(docx_path),
        "pdf": str(pdf_path) if pdf_path else None,
        "pages": pages,
        "within_limit": pages is not None and pages <= max_pages,
        "last_page_fill": last_page_fill(pdf_path) if pdf_path else None,
        "autofit_steps": steps_applied,
        "final_typography": {"body_pt": params["size_body"], "margins_in": params["margins_in"],
                             "spacing": params["spacing"], "font": params["font_body"]},
    }
    if pdf_path and (previews or first_page_preview):
        result["previews"] = render_previews(pdf_path, out_dir / f"{basename}_preview",
                                             dpi=80, first_page_only=first_page_preview and not previews)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cv", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--basename", required=True)
    parser.add_argument("--template", default="mirror",
                        help="mirror|classic|modern|compact, or a comma list to render options")
    parser.add_argument("--style")
    parser.add_argument("--max-pages", type=int, default=2)
    parser.add_argument("--autofit", action="store_true")
    parser.add_argument("--previews", action="store_true", help="PNG preview of every page")
    args = parser.parse_args()

    if Document is None:
        fail("python-docx is not installed. Run: pip install python-docx")
    templates = [t.strip() for t in args.template.split(",") if t.strip()]
    for t in templates:
        if t != "mirror" and t not in TEMPLATES:
            fail(f"Unknown template '{t}'", allowed=["mirror", *TEMPLATES])
    cv = load_json(args.cv)
    style = load_json(args.style) if args.style else None
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    multiple = len(templates) > 1
    results = []
    for t in templates:
        name = f"{args.basename}_{t}" if multiple else args.basename
        results.append(render_one(cv, t, style, out_dir, name, args.max_pages, args.autofit,
                                  args.previews, first_page_preview=multiple))
    warnings = []
    if not find_soffice():
        warnings.append("LibreOffice not found: DOCX only, page count unverified. Run scripts/doctor.py.")
    if style and template_notes(style):
        warnings.extend(template_notes(style))
    emit({"ok": all(r["within_limit"] or r["pages"] is None for r in results),
          "max_pages": args.max_pages, "results": results, "warnings": warnings})


def template_notes(style: dict) -> list[str]:
    return list(style.get("notes", []))


if __name__ == "__main__":
    main()
